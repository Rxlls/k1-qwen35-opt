#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""format_docx.py — 按赛题提交规范整理 Word 文档

1. 正文（含表格内文字）统一字体：宋体
2. 各级标题统一字体：黑体
3. 所有表格：自动调整方式 = 根据窗口调整表格（tblW=pct 5000 + tblLayout=autofit）
4. 代码块保留等宽字体（Consolas），避免代码变形

用法: python format_docx.py <in.docx> [out.docx]
"""
import sys
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY = "宋体"
HEAD = "黑体"
CODE = "Consolas"


def set_run_font(run, name):
    """同时设置西文与东亚字体（只设 font.name 不会作用于中文）"""
    if not run.text or not run.text.strip():
        return
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def set_style_font(doc, style_name, font):
    try:
        st = doc.styles[style_name]
    except KeyError:
        return False
    st.font.name = font
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font)
    return True


def fit_table_to_window(table):
    """tblW=pct 5000 (100% 页宽) + tblLayout=autofit"""
    tblPr = table._tbl.tblPr

    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        ts = tblPr.find(qn("w:tblStyle"))
        if ts is not None:
            ts.addnext(tblW)
        else:
            tblPr.insert(0, tblW)
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")

    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblW.addnext(layout)
    layout.set(qn("w:type"), "autofit")

    table.autofit = True


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "docs/REPORT_formatted.docx"
    dst = sys.argv[2] if len(sys.argv) > 2 else src

    doc = docx.Document(src)

    # 1) 样式层面：Normal 设为宋体（兜底，未显式设置的文字也生效）
    set_style_font(doc, "Normal", BODY)

    # 2) 段落层面
    n_head = n_body = n_code = 0
    for p in doc.paragraphs:
        sname = p.style.name or ""
        if sname.startswith("Heading") or sname in ("Title", "Subtitle"):
            font, n_head = HEAD, n_head + 1
        elif sname in ("Source Code", "Verbatim Char"):
            font, n_code = CODE, n_code + 1
        else:
            font, n_body = BODY, n_body + 1
        for r in p.runs:
            set_run_font(r, font)

    # 3) 表格内文字
    n_cell = 0
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_run_font(r, BODY)
                    n_cell += 1

    # 4) 表格自适应窗口
    n_tbl = 0
    for t in doc.tables:
        fit_table_to_window(t)
        n_tbl += 1

    doc.save(dst)

    print(f"保存: {dst}")
    print(f"  段落: 标题 {n_head} / 正文 {n_body} / 代码块 {n_code}")
    print(f"  表格: {n_tbl} 个（已设为根据窗口调整），单元格段落 {n_cell} 个")
    print(f"  字体: 正文={BODY}  标题={HEAD}  代码={CODE}")


if __name__ == "__main__":
    main()
